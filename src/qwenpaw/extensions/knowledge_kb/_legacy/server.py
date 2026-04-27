#!/usr/bin/env python3
from email.parser import BytesParser
from email.policy import default as email_policy
from html.parser import HTMLParser
from difflib import SequenceMatcher
import json
import os
import re
import math
import sqlite3
import sys
import threading
import uuid
from collections import Counter
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from xml.etree import ElementTree as ET
import zipfile
import unicodedata

import array
import llm_provider
import embedding_provider

try:
    from pypdf import PdfReader  # type: ignore
except Exception:  # pragma: no cover
    PdfReader = None

try:
    from PIL import Image  # type: ignore
except Exception:  # pragma: no cover
    Image = None

try:
    import pytesseract  # type: ignore
except Exception:  # pragma: no cover
    pytesseract = None

try:
    from docx import Document as DocxDocument  # type: ignore
    from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError  # type: ignore
except Exception:  # pragma: no cover
    DocxDocument = None
    class DocxPackageNotFoundError(Exception):  # type: ignore
        pass


BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent.parent
DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
DB_PATH = DATA_DIR / "knowledge.db"
DEMO_HTML = BASE_DIR / "demo.html"
BUILTIN_KB_DIR = BASE_DIR / "builtin_kb"

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_MANUAL_ENTRY_BYTES = 50 * 1024
ORPHAN_JOB_THRESHOLD_SECONDS = 600

LLM_ENABLED = bool(os.environ.get("DEEPSEEK_API_KEY"))
LLM_DEFAULT_PROVIDER = "deepseek"

EMBEDDING_KEY_CONFIGURED = bool(os.environ.get("DASHSCOPE_API_KEY"))
# env override: EMBEDDING_ENABLED=false|0|off|no → force off even with key present.
#               EMBEDDING_ENABLED=auto (default) → on iff key configured
#               EMBEDDING_ENABLED=true|1|on|yes → on iff key configured (else stays off)
_embed_env_raw = (os.environ.get("EMBEDDING_ENABLED", "auto") or "auto").strip().lower()
EMBEDDING_ENV_FORCED_OFF = _embed_env_raw in ("false", "0", "off", "no", "disabled")
# Runtime mutable — toggled via /api/v1/admin/embedding/toggle without restart.
EMBEDDING_ENABLED = EMBEDDING_KEY_CONFIGURED and not EMBEDDING_ENV_FORCED_OFF
EMBEDDING_DEFAULT_PROVIDER = "dashscope"
EMBEDDING_BATCH_SIZE = 10
HYBRID_ALPHA = 0.5


def embedding_status() -> dict:
    return {
        "enabled": EMBEDDING_ENABLED,
        "key_configured": EMBEDDING_KEY_CONFIGURED,
        "env_forced_off": EMBEDDING_ENV_FORCED_OFF,
        "provider": EMBEDDING_DEFAULT_PROVIDER,
    }


def toggle_embedding(on: bool) -> tuple[bool, str]:
    """Flip runtime embedding state. Returns (success, reason_if_rejected)."""
    global EMBEDDING_ENABLED
    if on:
        if not EMBEDDING_KEY_CONFIGURED:
            return False, "no_key"
        if EMBEDDING_ENV_FORCED_OFF:
            return False, "env_forced_off"
    EMBEDDING_ENABLED = bool(on)
    print(f"[embed] runtime toggled → enabled={EMBEDDING_ENABLED}", flush=True)
    return True, "ok"
RAG_SYSTEM_PROMPT = (
    "你是运维知识库的 RAG 合成助手。只用下面给出的「资料」回答用户问题。\n"
    "规则:\n"
    "1. 事实必须严格出自资料,禁止编造或补充资料外的信息\n"
    "2. 每条结论末尾用 [序号] 标注引用的资料编号,可多个如 [1][3]\n"
    "3. 如果资料不足以回答,直接说 \"资料不足以回答\",别硬答\n"
    "4. 用简洁中文,能用列表就列表\n"
    "5. 不要元注释,不要 \"根据资料...\",直接给答案"
)
LLM_FALLBACK_SYSTEM_PROMPT = (
    "你是运维知识库的辅助回答助手。用户的问题没有命中本地知识库。\n"
    "请基于通用知识给一个简洁、带边界的回答。\n"
    "如果问题太宽泛或信息不足，直接说明需要更多上下文，不要编造细节。\n"
    "不要给出根因结论。列出建议时每条不超过 80 字。"
)

for path in (DATA_DIR, UPLOAD_DIR, BUILTIN_KB_DIR):
    path.mkdir(parents=True, exist_ok=True)


SCHEMA = """
CREATE TABLE IF NOT EXISTS source_record (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    filename TEXT NOT NULL,
    source_type TEXT NOT NULL,
    source_scope TEXT DEFAULT 'tenant_private',
    builtin_pack_id TEXT,
    builtin_pack_version TEXT,
    mime_type TEXT,
    storage_path TEXT NOT NULL,
    uploaded_at TEXT NOT NULL,
    extracted_text_length INTEGER DEFAULT 0,
    note TEXT,
    archived_at TEXT,
    archive_reason TEXT,
    meta_json TEXT
);

CREATE TABLE IF NOT EXISTS ingestion_job (
    id TEXT PRIMARY KEY,
    filename TEXT NOT NULL,
    source_type TEXT NOT NULL,
    status TEXT NOT NULL,
    created_at TEXT NOT NULL,
    finished_at TEXT,
    unit_count INTEGER DEFAULT 0,
    note TEXT
);

CREATE TABLE IF NOT EXISTS knowledge_unit (
    id TEXT PRIMARY KEY,
    source_record_id INTEGER NOT NULL,
    chunk_index INTEGER NOT NULL,
    title TEXT NOT NULL,
    content TEXT NOT NULL,
    locator TEXT,
    source_type TEXT NOT NULL,
    source_scope TEXT DEFAULT 'tenant_private',
    created_at TEXT NOT NULL,
    archived_at TEXT,
    archive_reason TEXT,
    meta_json TEXT,
    FOREIGN KEY(source_record_id) REFERENCES source_record(id)
);
"""


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def get_conn():
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_conn()
    try:
        conn.execute("PRAGMA journal_mode=WAL")
        conn.executescript(SCHEMA)
        ensure_column(conn, "source_record", "source_scope", "TEXT DEFAULT 'tenant_private'")
        ensure_column(conn, "source_record", "builtin_pack_id", "TEXT")
        ensure_column(conn, "source_record", "builtin_pack_version", "TEXT")
        ensure_column(conn, "source_record", "archived_at", "TEXT")
        ensure_column(conn, "source_record", "archive_reason", "TEXT")
        ensure_column(conn, "knowledge_unit", "source_scope", "TEXT DEFAULT 'tenant_private'")
        ensure_column(conn, "knowledge_unit", "archived_at", "TEXT")
        ensure_column(conn, "knowledge_unit", "archive_reason", "TEXT")
        ensure_column(conn, "ingestion_job", "progress_pct", "INTEGER DEFAULT 0")
        ensure_column(conn, "ingestion_job", "current_stage", "TEXT")
        ensure_column(conn, "ingestion_job", "updated_at", "TEXT")
        ensure_column(conn, "knowledge_unit", "embedding", "BLOB")
        ensure_column(conn, "knowledge_unit", "embedding_model", "TEXT")
        ensure_column(conn, "knowledge_unit", "embedding_dim", "INTEGER")
        ensure_column(conn, "knowledge_unit", "embedding_at", "TEXT")
        conn.commit()
    finally:
        conn.close()


def ensure_column(conn: sqlite3.Connection, table: str, column: str, definition: str):
    columns = {row["name"] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}
    if column in columns:
        return
    conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def insert_curated_knowledge(
    *,
    title: str,
    content: str,
    meta: dict,
    source_scope: str = "runtime_curated",
    source_type: str = "document",
    filename: str | None = None,
    note: str | None = None,
) -> dict:
    created_at = now_iso()
    safe_name = filename or f"curated-{uuid.uuid4().hex[:8]}.txt"
    conn = get_conn()
    try:
        cur = conn.execute(
            """
            INSERT INTO source_record (
                filename, source_type, source_scope, builtin_pack_id, builtin_pack_version,
                mime_type, storage_path, uploaded_at, extracted_text_length, note, meta_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                safe_name,
                source_type,
                source_scope,
                None,
                None,
                "text/plain",
                "",
                created_at,
                len(content),
                note or "运行时沉淀",
                json.dumps(meta, ensure_ascii=False),
            ),
        )
        source_record_id = int(cur.lastrowid)
        ku_id = f"ku_{uuid.uuid4().hex[:12]}"
        conn.execute(
            """
            INSERT INTO knowledge_unit (
                id, source_record_id, chunk_index, title, content, locator,
                source_type, source_scope, created_at, meta_json
            ) VALUES (?, ?, 0, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                ku_id,
                source_record_id,
                title,
                content,
                "运行时沉淀",
                source_type,
                source_scope,
                created_at,
                json.dumps(meta, ensure_ascii=False),
            ),
        )
        conn.commit()
        return {
            "source_record_id": source_record_id,
            "knowledge_unit_id": ku_id,
            "created_at": created_at,
            "source_scope": source_scope,
        }
    finally:
        conn.close()


INGEST_STAGE_PCT = {
    "queued": 0,
    "extracting": 15,
    "chunking": 55,
    "indexing": 75,
    "embedding": 90,
    "success": 100,
    "failed": 0,
}


def _vector_to_blob(vec) -> bytes:
    return array.array("f", vec).tobytes()


def _blob_to_vector(blob) -> list[float]:
    if not blob:
        return []
    arr = array.array("f")
    arr.frombytes(blob)
    return list(arr)


def compute_unit_embeddings(ku_ids: list[str], batch_size: int = EMBEDDING_BATCH_SIZE) -> int:
    """Compute and store embeddings for the given units. Returns count written.

    Silently no-op if EMBEDDING_ENABLED is False. Logs and continues on batch
    failures so a transient API error doesn't kill a whole ingest.
    """
    if not EMBEDDING_ENABLED or not ku_ids:
        return 0

    conn = get_conn()
    try:
        placeholders = ",".join(["?"] * len(ku_ids))
        rows = conn.execute(
            f"SELECT id, title, content FROM knowledge_unit WHERE id IN ({placeholders})",
            list(ku_ids),
        ).fetchall()
    finally:
        conn.close()

    updated = 0
    for i in range(0, len(rows), batch_size):
        batch = rows[i:i + batch_size]
        texts = []
        for r in batch:
            title = (r["title"] or "").strip()
            content = (r["content"] or "").strip()
            composite = f"{title}\n\n{content}" if title else content
            # DashScope v4 max input is generous (~8K tokens); truncate defensively
            texts.append(composite[:6000])

        try:
            vectors = embedding_provider.embed_texts(
                texts,
                provider=EMBEDDING_DEFAULT_PROVIDER,
                batch_id=f"ingest-{i}-{len(batch)}",
            )
        except embedding_provider.EmbeddingError as exc:
            print(
                f"WARN:[embed] batch {i} failed ({type(exc).__name__}): {exc}",
                flush=True,
            )
            continue

        conn = get_conn()
        try:
            for r, vec in zip(batch, vectors):
                blob = _vector_to_blob(vec)
                conn.execute(
                    "UPDATE knowledge_unit SET embedding=?, embedding_model=?, embedding_dim=?, embedding_at=? WHERE id=?",
                    (
                        blob,
                        embedding_provider.DEFAULT_DASHSCOPE_MODEL,
                        len(vec),
                        now_iso(),
                        r["id"],
                    ),
                )
                updated += 1
            conn.commit()
        finally:
            conn.close()

    return updated


def _update_job(job_id: str, *, stage: str | None = None, pct: int | None = None, status: str | None = None, **extras) -> None:
    sets: list[str] = ["updated_at=?"]
    args: list[object] = [now_iso()]
    if stage is not None:
        sets.append("current_stage=?")
        args.append(stage)
    if pct is not None:
        sets.append("progress_pct=?")
        args.append(pct)
    if status is not None:
        sets.append("status=?")
        args.append(status)
    for k, v in extras.items():
        sets.append(f"{k}=?")
        args.append(v)
    args.append(job_id)
    conn = get_conn()
    try:
        conn.execute(f"UPDATE ingestion_job SET {', '.join(sets)} WHERE id=?", args)
        conn.commit()
    finally:
        conn.close()


def _run_ingest_worker(job_id: str, storage_path: str, filename: str, source_type: str, mime_type: str | None) -> None:
    started = time.monotonic() if False else None  # keep pure for now
    try:
        _update_job(job_id, stage="extracting", pct=INGEST_STAGE_PCT["extracting"], status="running")
        print(
            f"[ingest] {json.dumps({'ts': now_iso(), 'job_id': job_id, 'stage': 'extracting', 'filename': filename}, ensure_ascii=False)}",
            flush=True,
        )
        path = Path(storage_path)
        units, extracted_len, note, meta = extract_chunks(path, filename, source_type, mime_type)

        _update_job(job_id, stage="chunking", pct=INGEST_STAGE_PCT["chunking"])
        _update_job(job_id, stage="indexing", pct=INGEST_STAGE_PCT["indexing"])

        conn = get_conn()
        try:
            cur = conn.execute(
                """
                INSERT INTO source_record (
                    filename, source_type, source_scope, builtin_pack_id, builtin_pack_version,
                    mime_type, storage_path, uploaded_at, extracted_text_length, note, meta_json
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    filename,
                    source_type,
                    "tenant_private",
                    None,
                    None,
                    mime_type or "",
                    storage_path,
                    now_iso(),
                    extracted_len,
                    note,
                    json.dumps(meta, ensure_ascii=False),
                ),
            )
            source_record_id = int(cur.lastrowid)

            inserted_ids: list[str] = []
            for idx, unit in enumerate(units):
                ku_id = f"ku_{uuid.uuid4().hex[:12]}"
                inserted_ids.append(ku_id)
                conn.execute(
                    """
                    INSERT INTO knowledge_unit (
                        id, source_record_id, chunk_index, title, content, locator,
                        source_type, source_scope, created_at, meta_json
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        ku_id,
                        source_record_id,
                        idx,
                        unit["title"],
                        unit["content"],
                        unit["locator"],
                        source_type,
                        "tenant_private",
                        now_iso(),
                        json.dumps(unit.get("meta", {}), ensure_ascii=False),
                    ),
                )
            conn.commit()
        finally:
            conn.close()

        if EMBEDDING_ENABLED and inserted_ids:
            _update_job(job_id, stage="embedding", pct=INGEST_STAGE_PCT["embedding"])
            try:
                embed_count = compute_unit_embeddings(inserted_ids)
                print(
                    f"[ingest] {json.dumps({'ts': now_iso(), 'job_id': job_id, 'stage': 'embedded', 'embedded': embed_count, 'of': len(inserted_ids)}, ensure_ascii=False)}",
                    flush=True,
                )
            except Exception as exc:
                print(
                    f"WARN:[ingest] embedding phase failed for job {job_id}: {exc}",
                    flush=True,
                )

        _update_job(
            job_id,
            stage="success",
            pct=INGEST_STAGE_PCT["success"],
            status="success",
            finished_at=now_iso(),
            unit_count=len(units),
            note=note,
        )
        print(
            f"[ingest] {json.dumps({'ts': now_iso(), 'job_id': job_id, 'stage': 'success', 'unit_count': len(units), 'source_record_id': source_record_id}, ensure_ascii=False)}",
            flush=True,
        )
    except Exception as exc:
        _update_job(
            job_id,
            stage="failed",
            status="failed",
            finished_at=now_iso(),
            note=str(exc),
        )
        print(
            f"[ingest] {json.dumps({'ts': now_iso(), 'job_id': job_id, 'stage': 'failed', 'error': str(exc)}, ensure_ascii=False)}",
            flush=True,
        )


def reap_orphan_jobs(threshold_seconds: int = ORPHAN_JOB_THRESHOLD_SECONDS) -> int:
    cutoff = (
        datetime.now(timezone.utc) - timedelta(seconds=threshold_seconds)
    ).replace(microsecond=0).isoformat()
    conn = get_conn()
    try:
        cur = conn.execute(
            "UPDATE ingestion_job SET status=?, finished_at=?, note=? "
            "WHERE status='running' AND COALESCE(updated_at, created_at) < ?",
            ("failed", now_iso(), "启动时发现孤儿 job（进程重启前未完成）", cutoff),
        )
        conn.commit()
        return cur.rowcount
    finally:
        conn.close()


def parse_bool(value: str | None, default: bool = False) -> bool:
    if value is None:
        return default
    return value.strip().lower() in ("1", "true", "yes", "y", "on")


def parse_uploaded_file(handler: BaseHTTPRequestHandler) -> dict | None:
    content_type = handler.headers.get("Content-Type", "")
    if "multipart/form-data" not in content_type:
        return None
    length = int(handler.headers.get("Content-Length", "0") or "0")
    if length > MAX_UPLOAD_BYTES:
        raise ValueError(f"upload exceeds max size {MAX_UPLOAD_BYTES} bytes (got {length})")
    raw = handler.rfile.read(length)
    if not raw:
        return None
    pseudo_message = (
        f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("utf-8")
        + raw
    )
    message = BytesParser(policy=email_policy).parsebytes(pseudo_message)
    for part in message.iter_parts():
        if part.get_param("name", header="content-disposition") != "file":
            continue
        filename = part.get_filename()
        if not filename:
            continue
        return {
            "filename": filename,
            "mime_type": part.get_content_type(),
            "raw": part.get_payload(decode=True) or b"",
        }
    return None


def detect_source_type(filename: str, mime_type: str | None) -> str:
    lowered = filename.lower()
    if lowered.endswith((".xlsx", ".xlsm")) or mime_type in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel.sheet.macroEnabled.12",
    ):
        return "spreadsheet"
    if lowered.endswith(".pdf") or mime_type == "application/pdf":
        return "pdf"
    if mime_type and mime_type.startswith("image/"):
        return "image"
    if lowered.endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg")):
        return "image"
    return "document"


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


class _HtmlTextExtractor(HTMLParser):
    _SKIP = {"script", "style", "noscript", "template"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip_depth > 0:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if self._skip_depth == 0:
            cleaned = data.strip()
            if cleaned:
                self.parts.append(cleaned)


def strip_html_to_text(html_text: str) -> str:
    parser = _HtmlTextExtractor()
    try:
        parser.feed(html_text)
    except Exception:
        return re.sub(r"<[^>]+>", " ", html_text)
    return "\n".join(parser.parts)


def split_text(text: str) -> list[str]:
    parts = [seg.strip() for seg in re.split(r"\n\s*\n", text) if seg.strip()]
    if not parts:
        text = text.strip()
        if not text:
            return []
        parts = [text]
    chunks = []
    for part in parts:
        if len(part) <= 380:
            chunks.append(part)
        else:
            for i in range(0, len(part), 340):
                chunks.append(part[i:i + 380])
    return chunks


MARKDOWN_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
DOCX_HEADING_STYLE_RE = re.compile(r"^(?:Heading|标题)\s*(\d+)$")


def split_markdown_with_hierarchy(text: str) -> list[dict]:
    """Parse markdown into chunks that carry a section_path (heading lineage)."""
    if not text or not text.strip():
        return []
    heading_stack: list[tuple[int, str]] = []
    current_body: list[str] = []
    chunks: list[dict] = []

    def flush():
        nonlocal current_body
        body = "\n".join(current_body).strip()
        current_body = []
        if not body:
            return
        path = [t for (_, t) in heading_stack]
        for piece in split_text(body):
            chunks.append({"content": piece, "section_path": list(path)})

    for line in text.split("\n"):
        m = MARKDOWN_HEADING_RE.match(line)
        if m:
            flush()
            level = len(m.group(1))
            title = m.group(2).strip()
            heading_stack = [(lv, t) for (lv, t) in heading_stack if lv < level]
            heading_stack.append((level, title))
        else:
            current_body.append(line)
    flush()

    if not chunks:
        for piece in split_text(text):
            chunks.append({"content": piece, "section_path": []})
    return chunks


def docx_zipbomb_check(
    path: Path,
    max_unc_total: int = 200 * 1024 * 1024,
    max_unc_single: int = 50 * 1024 * 1024,
) -> str | None:
    """Returns None if safe, or an error string if the archive is suspicious."""
    try:
        with zipfile.ZipFile(path) as zf:
            total = 0
            for info in zf.infolist():
                if info.file_size > max_unc_single:
                    return f"docx single-file uncompressed size {info.file_size} exceeds {max_unc_single}"
                total += info.file_size
                if total > max_unc_total:
                    return f"docx total uncompressed size {total} exceeds {max_unc_total}"
    except zipfile.BadZipFile as exc:
        return f"not a valid zip: {exc}"
    return None


def extract_docx(path: Path, filename: str) -> tuple[list[dict], int, str, dict]:
    if DocxDocument is None:
        units = [
            {
                "title": f"{filename} · Word 资料",
                "content": "docx 已上传，但当前环境未装 python-docx，只保留资料记录。",
                "locator": "整份 docx",
                "meta": {"format": "docx", "docx_extract": False, "lightweight_mode": True, "section_path": []},
            }
        ]
        return units, 0, "Word 已导入，但未启用正文抽取（python-docx 未装）", {"format": "docx", "docx_extract": False}

    bomb_err = docx_zipbomb_check(path)
    if bomb_err:
        units = [
            {
                "title": f"{filename} · Word 资料",
                "content": f"docx 已上传，但文件被拒绝：{bomb_err}",
                "locator": "整份 docx",
                "meta": {"format": "docx", "docx_extract": False, "zipbomb_rejected": True, "section_path": []},
            }
        ]
        return units, 0, "docx 已导入，但文件异常未解析", {"format": "docx", "docx_extract": False}

    try:
        doc = DocxDocument(str(path))
    except DocxPackageNotFoundError as exc:
        units = [
            {
                "title": f"{filename} · Word 资料",
                "content": f"docx 已上传，但打不开（可能已加密或损坏）：{exc}",
                "locator": "整份 docx",
                "meta": {"format": "docx", "docx_extract": False, "error": str(exc), "section_path": []},
            }
        ]
        return units, 0, "docx 已导入，但打开失败（可能加密）", {"format": "docx", "docx_extract": False}
    except Exception as exc:
        units = [
            {
                "title": f"{filename} · Word 资料",
                "content": f"docx 已上传，但解析失败：{exc}",
                "locator": "整份 docx",
                "meta": {"format": "docx", "docx_extract": False, "error": str(exc), "section_path": []},
            }
        ]
        return units, 0, "docx 已导入，但解析失败", {"format": "docx", "docx_extract": False}

    heading_stack: list[tuple[int, str]] = []
    current_body: list[str] = []
    chunks: list[dict] = []
    total_text_len = 0

    def flush(current_path: list[str]):
        nonlocal current_body
        body = "\n".join(p for p in current_body if p.strip())
        current_body = []
        if not body.strip():
            return
        path_str = " · ".join(current_path) if current_path else "正文"
        for idx, piece in enumerate(split_text(body)):
            chunks.append(
                {
                    "title": f"{filename} · {path_str} · 段 {idx + 1}",
                    "content": piece,
                    "locator": f"{path_str} / 段 {idx + 1}",
                    "meta": {
                        "format": "docx",
                        "section_path": list(current_path),
                        "docx_extract": True,
                    },
                }
            )

    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        total_text_len += len(text)
        style_name = (para.style.name or "") if para.style else ""
        m = DOCX_HEADING_STYLE_RE.match(style_name)
        if m:
            level = int(m.group(1))
            flush([t for (_, t) in heading_stack])
            heading_stack = [(lv, t) for (lv, t) in heading_stack if lv < level]
            heading_stack.append((level, text))
        else:
            current_body.append(text)
    flush([t for (_, t) in heading_stack])

    if not chunks:
        units = [
            {
                "title": f"{filename} · Word 资料",
                "content": "docx 已上传，但没有抽取到正文内容。",
                "locator": "整份 docx",
                "meta": {"format": "docx", "docx_extract": False, "empty_text": True, "section_path": []},
            }
        ]
        return units, 0, "docx 已导入，但未抽取到正文", {"format": "docx", "docx_extract": False}

    return chunks, total_text_len, "docx 正文已解析", {"format": "docx", "docx_extract": True}


def clean_display_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text or "")
    text = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", " ", text)
    text = text.replace("\ufeff", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def tokenize_query(text: str) -> list[str]:
    parts = re.findall(r"[A-Za-z0-9_\-]+|[\u4e00-\u9fff]+", text.lower())
    tokens: list[str] = []
    for part in parts:
        tokens.append(part)
        if re.search(r"[\u4e00-\u9fff]", part) and len(part) > 2:
            for size in (2, 3):
                if len(part) >= size:
                    tokens.extend(part[i:i + size] for i in range(len(part) - size + 1))
    # stable de-dup
    seen = set()
    ordered = []
    for token in tokens:
        if token and token not in seen:
            seen.add(token)
            ordered.append(token)
    return ordered


def build_source_label(filename: str, source_type: str) -> str:
    if source_type == "image":
        return f"图片：{filename}"
    if source_type == "pdf":
        return f"PDF：{filename}"
    return filename


def effective_source_label(source_record: dict, source_type: str) -> str:
    meta = source_record.get("meta") or {}
    display_title = (meta.get("display_title") or "").strip()
    if display_title:
        return display_title
    return build_source_label(source_record.get("filename", ""), source_type)


def scope_label(source_scope: str) -> str:
    return {
        "system_builtin": "平台内置知识",
        "runtime_curated": "运行时沉淀",
        "tenant_private": "企业内部经验",
    }.get(source_scope or "", "企业内部经验")


def normalize_response_source_type(source_type: str) -> str:
    if source_type in ("pdf", "document", "spreadsheet"):
        return "document"
    if source_type == "image":
        return "image"
    return source_type


def normalize_excerpt(text: str, max_len: int = 80) -> str:
    compact = re.sub(r"\s+", " ", clean_display_text(text or "")).strip()
    if len(compact) <= max_len:
        return compact
    return compact[: max_len - 1] + "…"


def extract_summary_points(text: str, limit: int = 3) -> list[str]:
    cleaned = clean_display_text(text)
    if not cleaned:
        return []
    pieces = re.split(r"[。\n；;！？!?]+", cleaned)
    points: list[str] = []
    seen = set()
    for piece in pieces:
        part = re.sub(r"\s+", " ", piece).strip(" ,:：-")
        if len(part) < 6:
            continue
        if part in seen:
            continue
        seen.add(part)
        points.append(part)
        if len(points) >= limit:
            break
    if not points:
        compact = normalize_excerpt(cleaned, 80)
        if compact:
            points.append(compact)
    return points


def normalize_summary_text(text: str) -> str:
    compact = re.sub(r"\s+", "", clean_display_text(text or "")).lower()
    compact = re.sub(r"[，。；：、“”‘’（）()【】《》…,.!?！？:;#\\-_/]", "", compact)
    return compact


def is_redundant_summary_pair(primary: str, secondary: str) -> bool:
    left = normalize_summary_text(primary)
    right = normalize_summary_text(secondary)
    if not left or not right:
        return False
    if left == right:
        return True
    if left in right or right in left:
        return True
    short_len = min(len(left), len(right))
    if short_len < 6:
        return False
    overlap = sum(1 for ch in set(right) if ch in set(left))
    ratio = overlap / max(1, len(set(right)))
    if ratio >= 0.8:
        return True
    return SequenceMatcher(a=left, b=right).ratio() >= 0.6


def pick_answer_points(evidence: list[dict], limit: int = 2) -> list[str]:
    candidates: list[tuple[int, str]] = []
    seen = set()
    for idx, ev in enumerate(evidence[:5]):
        points = extract_summary_points(ev.get("chunk_text") or ev.get("chunk_summary") or "", limit=3)
        for point in points:
            compact = re.sub(r"\s+", " ", point).strip()
            if not compact or compact in seen:
                continue
            seen.add(compact)
            penalty = 0
            if compact.startswith("#"):
                penalty += 80
            if len(compact) < 12:
                penalty += 30
            candidates.append((idx * 10 + penalty, compact))
    candidates.sort(key=lambda x: x[0])
    return [text for _, text in candidates[:limit]]


def build_summary_from_evidence(evidence: list[dict]) -> str:
    if not evidence:
        return "当前没有在知识库中找到足够相关的资料。"

    top = evidence[0]
    answer_points = pick_answer_points(evidence, limit=2)
    if answer_points:
        if len(answer_points) == 1:
            return answer_points[0]
        if is_redundant_summary_pair(answer_points[0], answer_points[1]):
            return answer_points[0]
        return f"{answer_points[0]}。补充信息：{answer_points[1]}"

    top_points = extract_summary_points(top.get("chunk_text") or top.get("chunk_summary") or "", limit=2)
    top_excerpt = "；".join(top_points) if top_points else normalize_excerpt(top.get("chunk_text") or top.get("chunk_summary") or "")
    return top_excerpt


def column_letters_to_index(ref: str) -> int:
    letters = "".join(ch for ch in ref if ch.isalpha())
    value = 0
    for ch in letters:
        value = value * 26 + (ord(ch.upper()) - 64)
    return max(value - 1, 0)


def extract_xlsx(path: Path, filename: str) -> tuple[list[dict], int, str, dict]:
    ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as zf:
        workbook_xml = ET.fromstring(zf.read("xl/workbook.xml"))
        rels_xml = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        rel_map = {
            rel.attrib["Id"]: rel.attrib["Target"]
            for rel in rels_xml
        }
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            shared_root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in shared_root.findall("a:si", ns):
                text = "".join(t.text or "" for t in si.findall(".//a:t", ns))
                shared_strings.append(text.strip())

        chunks: list[dict] = []
        total_text_len = 0
        sheets = workbook_xml.find("a:sheets", ns)
        if sheets is None:
            return [], 0, "Excel 已导入，但未找到工作表", {"format": "xlsx", "sheet_extract": False}

        for sheet in sheets.findall("a:sheet", ns):
            sheet_name = sheet.attrib.get("name", "Sheet")
            rel_id = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            target = rel_map.get(rel_id or "")
            if not target:
                continue
            # Normalize: some xlsx emit absolute OOXML refs like "/xl/worksheets/sheet1.xml"
            target_clean = target.lstrip("/")
            sheet_path = target_clean if target_clean.startswith("xl/") else f"xl/{target_clean}"
            if sheet_path not in zf.namelist():
                continue
            sheet_root = ET.fromstring(zf.read(sheet_path))
            rows = sheet_root.findall(".//a:sheetData/a:row", ns)
            parsed_rows: list[tuple[str, list[str]]] = []
            for row in rows:
                cell_values: list[str] = []
                for cell in row.findall("a:c", ns):
                    ref = cell.attrib.get("r", "")
                    cell_type = cell.attrib.get("t")
                    v = cell.find("a:v", ns)
                    raw_value = (v.text or "").strip() if v is not None else ""
                    if cell_type == "s" and raw_value.isdigit():
                        idx = int(raw_value)
                        value = clean_display_text(shared_strings[idx] if 0 <= idx < len(shared_strings) else "")
                    elif cell_type == "inlineStr":
                        value = clean_display_text("".join(t.text or "" for t in cell.findall(".//a:is//a:t", ns)))
                    elif cell_type == "str":
                        value = clean_display_text(raw_value)
                    else:
                        value = clean_display_text(raw_value)
                    if value:
                        col_idx = column_letters_to_index(ref)
                        while len(cell_values) < col_idx:
                            cell_values.append("")
                        if len(cell_values) == col_idx:
                            cell_values.append(value)
                        else:
                            cell_values[col_idx] = value
                normalized = [v.strip() for v in cell_values if v and v.strip()]
                if not normalized:
                    continue
                parsed_rows.append((row.attrib.get("r", "?"), normalized))

            if not parsed_rows:
                continue

            header_row_no, header_values = parsed_rows[0]
            data_rows = parsed_rows[1:] if len(parsed_rows) > 1 else parsed_rows
            batch_size = 50
            for start in range(0, len(data_rows), batch_size):
                group = data_rows[start:start + batch_size]
                lines = []
                row_ids = []
                for row_no, values in group:
                    row_ids.append(str(row_no))
                    if header_values and len(header_values) > 1:
                        pairs = []
                        for idx, value in enumerate(values):
                            key = header_values[idx] if idx < len(header_values) else f"列{idx + 1}"
                            pairs.append(f"{key}={value}")
                        line = "；".join(pairs)
                    else:
                        line = " | ".join(values)
                    lines.append(line)
                content = "\n".join(lines)
                total_text_len += len(content)
                row_range = f"{row_ids[0]}-{row_ids[-1]}" if len(row_ids) > 1 else row_ids[0]
                title = f"{filename} · {sheet_name} · 第 {row_range} 行"
                if header_values:
                    header_preview = " / ".join(header_values[:4])
                    title = f"{title} · {header_preview}"
                chunks.append(
                    {
                        "title": title,
                        "content": content,
                        "locator": f"{sheet_name} / 第 {row_range} 行",
                        "meta": {"format": "xlsx", "sheet": sheet_name, "header_row": header_row_no},
                    }
                )

    if not chunks:
        return [], 0, "Excel 已导入，但没有抽取到有效单元格内容", {"format": "xlsx", "sheet_extract": False}
    return chunks, total_text_len, "Excel 内容已解析", {"format": "xlsx", "sheet_extract": True}


def list_builtin_manifests() -> list[dict]:
    packs: list[dict] = []
    if not BUILTIN_KB_DIR.exists():
        return packs
    for manifest_path in sorted(BUILTIN_KB_DIR.rglob("manifest.json")):
        payload = json.loads(manifest_path.read_text(encoding="utf-8"))
        payload["manifest_path"] = str(manifest_path)
        payload["declared_file_count"] = len(payload.get("files", []))
        packs.append(payload)
    return packs


def list_builtin_pack_stats() -> list[dict]:
    conn = get_conn()
    try:
        rows = conn.execute(
            """
            SELECT sr.builtin_pack_id,
                   sr.builtin_pack_version,
                   COUNT(DISTINCT sr.id) AS source_count,
                   COUNT(ku.id) AS unit_count,
                   MAX(sr.uploaded_at) AS imported_at
            FROM source_record sr
            LEFT JOIN knowledge_unit ku ON ku.source_record_id = sr.id
            WHERE sr.source_scope = 'system_builtin'
              AND sr.builtin_pack_id IS NOT NULL
              AND sr.archived_at IS NULL
              AND (ku.archived_at IS NULL OR ku.archived_at = '')
            GROUP BY sr.builtin_pack_id, sr.builtin_pack_version
            ORDER BY imported_at DESC
            """
        ).fetchall()
        return [dict(row) for row in rows]
    finally:
        conn.close()


def get_builtin_source_record(pack_id: str, version: str, filename: str) -> dict | None:
    conn = get_conn()
    try:
        row = conn.execute(
            """
            SELECT *
            FROM source_record
            WHERE source_scope = 'system_builtin'
              AND builtin_pack_id = ?
              AND builtin_pack_version = ?
              AND filename = ?
              AND archived_at IS NULL
            LIMIT 1
            """,
            (pack_id, version, filename),
        ).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def delete_builtin_pack(pack_id: str):
    conn = get_conn()
    try:
        rows = conn.execute(
            """
            SELECT id
            FROM source_record
            WHERE source_scope = 'system_builtin'
              AND builtin_pack_id = ?
            """,
            (pack_id,),
        ).fetchall()
        source_ids = [row["id"] for row in rows]
        if not source_ids:
            return
        placeholders = ",".join(["?"] * len(source_ids))
        conn.execute(
            f"DELETE FROM knowledge_unit WHERE source_record_id IN ({placeholders})",
            source_ids,
        )
        conn.execute(
            f"DELETE FROM source_record WHERE id IN ({placeholders})",
            source_ids,
        )
        conn.commit()
    finally:
        conn.close()


def list_builtin_pack_status() -> list[dict]:
    stats_map = {
        (row["builtin_pack_id"], row["builtin_pack_version"]): row
        for row in list_builtin_pack_stats()
    }
    items: list[dict] = []
    for pack in list_builtin_manifests():
        stat = stats_map.get((pack["pack_id"], pack["version"]), {})
        items.append(
            {
                "pack_id": pack["pack_id"],
                "version": pack["version"],
                "title": pack.get("title") or pack["pack_id"],
                "description": pack.get("description", ""),
                "scope_label": pack.get("scope_label", "平台内置知识"),
                "enabled": bool(pack.get("enabled", True)),
                "manifest_path": pack["manifest_path"],
                "declared_file_count": pack["declared_file_count"],
                "imported_source_count": int(stat.get("source_count", 0) or 0),
                "imported_unit_count": int(stat.get("unit_count", 0) or 0),
                "imported_at": stat.get("imported_at"),
            }
        )
    return items


def list_source_records(
    limit: int = 50,
    offset: int = 0,
    filters: dict | None = None,
    include_archived: bool = False,
) -> dict:
    filters = normalize_filters(filters)
    where = []
    params: list[object] = []
    if not include_archived:
        where.append("sr.archived_at IS NULL")
    if filters.get("source_scope"):
        where.append("sr.source_scope = ?")
        params.append(filters["source_scope"])
    if filters.get("source_type"):
        where.append("sr.source_type = ?")
        params.append(filters["source_type"])
    if filters.get("builtin_pack_id"):
        where.append("sr.builtin_pack_id = ?")
        params.append(filters["builtin_pack_id"])
    if filters.get("filename"):
        where.append("sr.filename LIKE ?")
        params.append(f"%{filters['filename']}%")
    where_clause = f"WHERE {' AND '.join(where)}" if where else ""

    conn = get_conn()
    try:
        total_row = conn.execute(
            f"""
            SELECT COUNT(*) AS total
            FROM source_record sr
            {where_clause}
            """,
            params,
        ).fetchone()
        rows = conn.execute(
            f"""
            SELECT sr.id,
                   sr.filename,
                   sr.source_type,
                   sr.source_scope,
                   sr.builtin_pack_id,
                   sr.builtin_pack_version,
                   sr.uploaded_at,
                   sr.note,
                   sr.archived_at,
                   sr.archive_reason,
                   sr.meta_json,
                   COUNT(ku.id) AS unit_count
            FROM source_record sr
            LEFT JOIN knowledge_unit ku ON ku.source_record_id = sr.id
            {where_clause}
            GROUP BY sr.id
            ORDER BY COALESCE(sr.archived_at, sr.uploaded_at) DESC, sr.id DESC
            LIMIT ? OFFSET ?
            """,
            [*params, limit, offset],
        ).fetchall()
        items = []
        for row in rows:
            item = dict(row)
            item["meta"] = json.loads(item.pop("meta_json") or "{}")
            items.append(item)
        return {
            "items": items,
            "total": int(total_row["total"] if total_row else 0),
            "limit": limit,
            "offset": offset,
        }
    finally:
        conn.close()


def get_source_record_detail(source_record_id: int, include_archived: bool = False) -> dict | None:
    conn = get_conn()
    try:
        where = "WHERE sr.id = ?"
        if not include_archived:
            where += " AND (sr.archived_at IS NULL)"
        row = conn.execute(
            f"""
            SELECT sr.*
            FROM source_record sr
            {where}
            """,
            (source_record_id,),
        ).fetchone()
        if not row:
            return None
        record = dict(row)
        record["meta"] = json.loads(record.pop("meta_json") or "{}")
        units = conn.execute(
            """
            SELECT ku.id, ku.title, ku.content, ku.locator, ku.source_type, ku.source_scope, ku.created_at, ku.meta_json
            FROM knowledge_unit ku
            WHERE ku.source_record_id = ?
            ORDER BY ku.chunk_index ASC
            """,
            (source_record_id,),
        ).fetchall()
        unit_items = []
        for unit in units:
            item = dict(unit)
            item["meta"] = json.loads(item.pop("meta_json") or "{}")
            unit_items.append(item)
        record["units"] = unit_items
        return record
    finally:
        conn.close()


def update_source_record_metadata(
    source_record_id: int,
    display_title: str | None = None,
    tags: list[str] | None = None,
    note: str | None = None,
    source_scope: str | None = None,
) -> dict:
    conn = get_conn()
    try:
        row = conn.execute(
            "SELECT * FROM source_record WHERE id = ?",
            (source_record_id,),
        ).fetchone()
        if not row:
            raise ValueError("source record not found")
        record = dict(row)
        meta = json.loads(record.get("meta_json") or "{}")
        if display_title is not None:
            display_title = display_title.strip()
            if display_title:
                meta["display_title"] = display_title
            else:
                meta.pop("display_title", None)
        if tags is not None:
            normalized_tags = [t.strip() for t in tags if t and t.strip()]
            if normalized_tags:
                meta["tags"] = normalized_tags
            else:
                meta.pop("tags", None)
        next_scope = record.get("source_scope") or "tenant_private"
        if source_scope is not None and source_scope.strip():
            requested_scope = source_scope.strip()
            if record.get("builtin_pack_id") and requested_scope != "system_builtin":
                raise ValueError("builtin knowledge cannot change source_scope")
            next_scope = requested_scope
        next_note = note.strip() if note is not None else record.get("note", "")
        if next_scope == "system_builtin":
            meta["scope_label"] = "平台内置知识"
        elif next_scope == "runtime_curated":
            meta["scope_label"] = "运行时沉淀"
        else:
            meta["scope_label"] = "企业内部经验"

        conn.execute(
            """
            UPDATE source_record
            SET source_scope = ?, note = ?, meta_json = ?
            WHERE id = ?
            """,
            (next_scope, next_note, json.dumps(meta, ensure_ascii=False), source_record_id),
        )
        conn.execute(
            """
            UPDATE knowledge_unit
            SET source_scope = ?
            WHERE source_record_id = ?
            """,
            (next_scope, source_record_id),
        )
        conn.commit()
        return {
            "source_record_id": source_record_id,
            "source_scope": next_scope,
            "note": next_note,
            "meta": meta,
        }
    finally:
        conn.close()


def archive_source_records(source_record_ids: list[int], reason: str) -> dict:
    archived_at = now_iso()
    conn = get_conn()
    try:
        ids = [int(i) for i in source_record_ids if i is not None]
        if not ids:
            return {"updated_source_count": 0, "updated_unit_count": 0}
        placeholders = ",".join(["?"] * len(ids))
        cur_sr = conn.execute(
            f"""
            UPDATE source_record
            SET archived_at = ?, archive_reason = ?
            WHERE id IN ({placeholders})
            """,
            [archived_at, reason, *ids],
        )
        cur_ku = conn.execute(
            f"""
            UPDATE knowledge_unit
            SET archived_at = ?, archive_reason = ?
            WHERE source_record_id IN ({placeholders})
            """,
            [archived_at, reason, *ids],
        )
        conn.commit()
        return {
            "updated_source_count": cur_sr.rowcount,
            "updated_unit_count": cur_ku.rowcount,
            "archived_at": archived_at,
            "reason": reason,
        }
    finally:
        conn.close()


def unarchive_source_records(source_record_ids: list[int]) -> dict:
    conn = get_conn()
    try:
        ids = [int(i) for i in source_record_ids if i is not None]
        if not ids:
            return {"updated_source_count": 0, "updated_unit_count": 0}
        placeholders = ",".join(["?"] * len(ids))
        cur_sr = conn.execute(
            f"""
            UPDATE source_record
            SET archived_at = NULL, archive_reason = NULL
            WHERE id IN ({placeholders})
            """,
            ids,
        )
        cur_ku = conn.execute(
            f"""
            UPDATE knowledge_unit
            SET archived_at = NULL, archive_reason = NULL
            WHERE source_record_id IN ({placeholders})
            """,
            ids,
        )
        conn.commit()
        return {
            "updated_source_count": cur_sr.rowcount,
            "updated_unit_count": cur_ku.rowcount,
        }
    finally:
        conn.close()


def import_enabled_builtin_packs(force: bool = False, pack_id: str | None = None) -> dict:
    packs = [pack for pack in list_builtin_manifests() if pack.get("enabled", True)]
    if pack_id:
        packs = [pack for pack in packs if pack["pack_id"] == pack_id]

    imported: list[dict] = []
    skipped: list[str] = []
    replaced: list[str] = []

    stats = list_builtin_pack_stats()
    exact_stats = {(row["builtin_pack_id"], row["builtin_pack_version"]): row for row in stats}
    by_pack: dict[str, list[dict]] = {}
    for row in stats:
        by_pack.setdefault(row["builtin_pack_id"], []).append(row)

    for pack in packs:
        prior_versions = by_pack.get(pack["pack_id"], [])
        exact = exact_stats.get((pack["pack_id"], pack["version"]))
        should_replace = force or any(row["builtin_pack_version"] != pack["version"] for row in prior_versions)
        if should_replace and prior_versions:
            delete_builtin_pack(pack["pack_id"])
            replaced.append(pack["pack_id"])
        elif exact and int(exact.get("source_count", 0) or 0) >= len(pack.get("files", [])):
            skipped.append(pack["pack_id"])
            continue

        manifest_dir = Path(pack["manifest_path"]).parent
        created_at = now_iso()
        conn = get_conn()
        try:
            imported_files = 0
            imported_units = 0
            for file_cfg in pack.get("files", []):
                file_path = manifest_dir / file_cfg["path"]
                if not file_path.exists():
                    raise FileNotFoundError(f"builtin knowledge file not found: {file_path}")
                if not should_replace and get_builtin_source_record(pack["pack_id"], pack["version"], file_path.name):
                    continue

                file_source_type = detect_source_type(file_path.name, None)
                units, extracted_len, note, meta = extract_chunks(file_path, file_path.name, file_source_type, None)
                cur = conn.execute(
                    """
                    INSERT INTO source_record (
                        filename, source_type, source_scope, builtin_pack_id, builtin_pack_version,
                        mime_type, storage_path, uploaded_at, extracted_text_length, note, meta_json
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        file_path.name,
                        file_source_type,
                        "system_builtin",
                        pack["pack_id"],
                        pack["version"],
                        "",
                        str(file_path),
                        created_at,
                        extracted_len,
                        note,
                        json.dumps(
                            {
                                **meta,
                                "scope_label": pack.get("scope_label", "平台内置知识"),
                                "builtin_pack_title": pack.get("title") or pack["pack_id"],
                                "declared_title": file_cfg.get("title", file_path.name),
                                "tags": list(file_cfg.get("tags", [])),
                            },
                            ensure_ascii=False,
                        ),
                    ),
                )
                source_record_id = int(cur.lastrowid)
                for idx, unit in enumerate(units):
                    conn.execute(
                        """
                        INSERT INTO knowledge_unit (id, source_record_id, chunk_index, title, content, locator, source_type, source_scope, created_at, meta_json)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            f"ku_{uuid.uuid4().hex[:12]}",
                            source_record_id,
                            idx,
                            unit["title"],
                            unit["content"],
                            unit["locator"],
                            file_source_type,
                            "system_builtin",
                            created_at,
                            json.dumps(
                                {
                                    **unit.get("meta", {}),
                                    "scope_label": pack.get("scope_label", "平台内置知识"),
                                    "builtin_pack_id": pack["pack_id"],
                                    "builtin_pack_version": pack["version"],
                                    "declared_title": file_cfg.get("title", file_path.name),
                                    "tags": list(file_cfg.get("tags", [])),
                                },
                                ensure_ascii=False,
                            ),
                        ),
                    )
                imported_files += 1
                imported_units += len(units)
            conn.commit()
        except Exception:
            conn.rollback()
            conn.close()
            raise
        finally:
            conn.close()

        imported.append(
            {
                "pack_id": pack["pack_id"],
                "version": pack["version"],
                "title": pack.get("title") or pack["pack_id"],
                "imported_source_count": imported_files,
                "imported_unit_count": imported_units,
            }
        )

    return {
        "requested_pack_id": pack_id,
        "force": force,
        "imported": imported,
        "skipped": skipped,
        "replaced": sorted(set(replaced)),
        "items": list_builtin_pack_status(),
    }


def normalize_filters(filters: dict | None) -> dict:
    raw = filters or {}
    normalized = {}
    for key in ("source_scope", "source_type", "builtin_pack_id", "filename"):
        value = raw.get(key)
        if value is None:
            continue
        if isinstance(value, str):
            value = value.strip()
        if value:
            normalized[key] = value
    return normalized


# BM25 parameters (Okapi BM25, standard defaults)
BM25_K1 = 1.5
BM25_B = 0.75
# Confidence mapping: map raw (BM25 + boosts) score to [0.48, 0.95].
# SAT_SCORE is the "fully confident" raw score — anything above saturates to 0.95.
BM25_SAT_SCORE = 8.0
CONFIDENCE_FLOOR = 0.48
CONFIDENCE_CEIL = 0.95


def query_units(query: str, limit: int = 3, filters: dict | None = None) -> tuple[list[dict], bool]:
    query_tokens = tokenize_query(query)
    lowered_query = query.lower().strip()
    filters = normalize_filters(filters)
    where = []
    params: list[object] = []
    where.append("sr.archived_at IS NULL")
    where.append("(ku.archived_at IS NULL OR ku.archived_at = '')")
    if filters.get("source_scope"):
        where.append("sr.source_scope = ?")
        params.append(filters["source_scope"])
    if filters.get("source_type"):
        where.append("ku.source_type = ?")
        params.append(filters["source_type"])
    if filters.get("builtin_pack_id"):
        where.append("sr.builtin_pack_id = ?")
        params.append(filters["builtin_pack_id"])
    where_clause = f"WHERE {' AND '.join(where)}" if where else ""
    conn = get_conn()
    try:
        rows = conn.execute(
            f"""
            SELECT ku.id, ku.source_type, ku.source_scope, ku.title, ku.content, ku.locator, ku.created_at,
                   ku.embedding, sr.filename, sr.uploaded_at, sr.builtin_pack_id, sr.meta_json
            FROM knowledge_unit ku
            JOIN source_record sr ON sr.id = ku.source_record_id
            {where_clause}
            ORDER BY ku.created_at DESC
            """,
            params,
        ).fetchall()
    finally:
        conn.close()

    if not rows:
        return [], True

    # Build corpus for BM25: tokenize each candidate's haystack once.
    docs: list[tuple[dict, dict, Counter, int, list]] = []
    for row in rows:
        item = dict(row)
        meta = json.loads(item.pop("meta_json") or "{}")
        display_title = (meta.get("display_title") or "").strip()
        haystack_text = f"{item['title']} {item['content']} {item['filename']} {display_title}"
        doc_tokens = tokenize_query(haystack_text)
        counter = Counter(doc_tokens)
        unit_vec = _blob_to_vector(item.pop("embedding", None))
        docs.append((item, meta, counter, len(doc_tokens), unit_vec))

    # Compute query embedding once (optional). Failure → pure BM25 fallback.
    query_vec: list[float] = []
    if EMBEDDING_ENABLED and query.strip() and any(d[4] for d in docs):
        try:
            q_vectors = embedding_provider.embed_texts(
                [query],
                provider=EMBEDDING_DEFAULT_PROVIDER,
                batch_id=f"q-{uuid.uuid4().hex[:8]}",
            )
            query_vec = q_vectors[0] if q_vectors else []
        except embedding_provider.EmbeddingError as exc:
            print(f"WARN:[query] embedding failed, BM25-only: {exc}", flush=True)
            query_vec = []

    n_docs = len(docs)
    total_len = sum(d[3] for d in docs)
    avgdl = (total_len / n_docs) if n_docs else 1.0
    if avgdl <= 0:
        avgdl = 1.0

    # Document frequency per query token (Okapi idf with +1 smoothing)
    idf: dict[str, float] = {}
    for qt in query_tokens:
        df = sum(1 for _, _, counter, _, _ in docs if qt in counter)
        idf[qt] = max(0.0, math.log((n_docs - df + 0.5) / (df + 0.5) + 1.0))

    scored = []
    for item, meta, counter, doc_len, unit_vec in docs:
        if doc_len == 0:
            continue
        bm25 = 0.0
        token_hits = 0
        for qt in query_tokens:
            tf = counter.get(qt, 0)
            if tf == 0:
                continue
            token_hits += 1
            numerator = tf * (BM25_K1 + 1.0)
            denom = tf + BM25_K1 * (1.0 - BM25_B + BM25_B * doc_len / avgdl)
            bm25 += idf[qt] * (numerator / denom)

        haystack_lower = (
            f"{item['title']} {item['content']} {item['filename']} "
            f"{(meta.get('display_title') or '')}"
        ).lower()
        # Identity fields: title + filename + display_title (excludes content body).
        identity_lower = (
            f"{item['title']} {item['filename']} {(meta.get('display_title') or '')}"
        ).lower()
        identity_bonus = 0.0
        for qt in query_tokens:
            if qt and qt in identity_lower:
                identity_bonus += idf[qt] * 0.8

        substring_bonus = 1.5 if (lowered_query and lowered_query in haystack_lower) else 0.0
        scope_bonus = 0.1 if item["source_scope"] != "system_builtin" else 0.0

        # Cosine similarity (semantic) — 0 if we have no query/unit embedding
        cos = 0.0
        if query_vec and unit_vec:
            cos = max(0.0, embedding_provider.cosine_sim(query_vec, unit_vec))

        # Skip docs with no signal in either modality
        if token_hits == 0 and substring_bonus == 0.0 and cos < 0.25:
            continue

        raw = bm25 + substring_bonus + scope_bonus + identity_bonus
        bm25_norm = min(1.0, raw / BM25_SAT_SCORE)

        # Hybrid: α·BM25 + (1-α)·cos, but only if we actually have a cosine signal
        if query_vec and unit_vec:
            final_norm = HYBRID_ALPHA * bm25_norm + (1.0 - HYBRID_ALPHA) * cos
        else:
            final_norm = bm25_norm

        confidence_score = min(
            CONFIDENCE_CEIL,
            CONFIDENCE_FLOOR + final_norm * (CONFIDENCE_CEIL - CONFIDENCE_FLOOR),
        )
        scored.append(
            {
                "evidence_id": item["id"],
                "confidence_score": round(confidence_score, 2),
                "confidence_level": "high" if confidence_score >= 0.78 else "medium" if confidence_score >= 0.62 else "low",
                "chunk_summary": clean_display_text(item["title"]),
                "chunk_text": clean_display_text(item["content"]),
                "source_time": item["uploaded_at"],
                "source_type": normalize_response_source_type(item["source_type"]),
                "citation": {
                    "source_label": effective_source_label({**item, "meta": meta}, item["source_type"]),
                    "source_type": normalize_response_source_type(item["source_type"]),
                    "source_scope_label": meta.get("scope_label", scope_label(item["source_scope"])),
                    "source_time": (item["uploaded_at"] or "")[:10],
                    "locator": item["locator"] or "-",
                },
                "meta": meta,
                "source_scope": item["source_scope"],
                "_raw_bm25": round(bm25, 3),
                "_cosine": round(cos, 3),
            }
        )

    scored.sort(key=lambda x: (x["confidence_score"], x["source_time"]), reverse=True)
    top = scored[:limit]
    sufficient = any((ev.get("confidence_score") or 0) >= 0.62 for ev in top)
    insufficient = len(top) == 0 or not sufficient

    algo = "hybrid" if query_vec else "bm25"
    print(
        f"[query] {json.dumps({'ts': now_iso(), 'query_hash': query[:40], 'candidates': n_docs, 'scored': len(scored), 'top_score': (scored[0]['confidence_score'] if scored else 0), 'sufficient': not insufficient, 'algo': algo}, ensure_ascii=False)}",
        flush=True,
    )

    # Strip internal debug fields before returning
    for ev in top:
        ev.pop("_raw_bm25", None)
        ev.pop("_cosine", None)
    return top, insufficient


def build_query_response(query: str, filters: dict | None = None) -> dict:
    filters = normalize_filters(filters)
    evidence, insufficient = query_units(query, filters=filters)
    filter_hint = ""
    if filters:
        filter_parts = []
        if filters.get("source_scope"):
            filter_parts.append(f"来源层级={scope_label(filters['source_scope'])}")
        if filters.get("source_type"):
            filter_parts.append(f"资料类型={filters['source_type']}")
        if filters.get("builtin_pack_id"):
            filter_parts.append(f"内置包={filters['builtin_pack_id']}")
        filter_hint = "（已按 " + " / ".join(filter_parts) + " 过滤）"
    if insufficient:
        return {
            "query_id": f"q_{uuid.uuid4().hex[:10]}",
            "answer_intent": "lookup",
            "layout_mode": "compact_chat",
            "summary": f"当前没有在知识库中找到足够相关的资料{filter_hint}。你可以先导入 Markdown、PDF 或图片，再重新提问。",
            "relevant_evidence": [],
            "troubleshooting_checklist": [],
            "possible_causes": [],
            "evidence_boundary_statement": f"证据不足：当前知识库里还没有和这个问题足够接近的资料{filter_hint}。",
            "flags": {
                "insufficient_evidence": True,
                "contradictory_evidence": False,
                "access_limited_evidence": False,
                "stale_evidence": False,
            },
        }

    return {
        "query_id": f"q_{uuid.uuid4().hex[:10]}",
        "answer_intent": "lookup",
        "layout_mode": "compact_chat",
        "summary": build_summary_from_evidence(evidence),
        "relevant_evidence": evidence,
        "troubleshooting_checklist": [],
        "possible_causes": [],
        "evidence_boundary_statement": "当前返回的是资料检索结果，不直接给排障步骤。你可以先打开最相关资料，再继续追问。",
        "flags": {
            "insufficient_evidence": False,
            "contradictory_evidence": False,
            "access_limited_evidence": False,
            "stale_evidence": False,
        },
    }


def list_source_summary() -> list[dict]:
    conn = get_conn()
    try:
        rows = conn.execute(
            """
            SELECT sr.source_scope,
                   ku.source_type,
                   sr.builtin_pack_id,
                   COUNT(ku.id) AS unit_count,
                   COUNT(DISTINCT sr.id) AS source_count,
                   MAX(ku.created_at) AS latest_created_at
            FROM knowledge_unit ku
            JOIN source_record sr ON sr.id = ku.source_record_id
            WHERE sr.archived_at IS NULL
              AND (ku.archived_at IS NULL OR ku.archived_at = '')
            GROUP BY sr.source_scope, ku.source_type, sr.builtin_pack_id
            ORDER BY unit_count DESC, latest_created_at DESC
            """
        ).fetchall()
        return [dict(row) for row in rows]
    finally:
        conn.close()


def get_unit_detail(unit_id: str) -> dict | None:
    conn = get_conn()
    try:
        row = conn.execute(
            """
            SELECT ku.id, ku.source_type, ku.source_scope, ku.title, ku.content, ku.locator, ku.created_at,
                   sr.id AS source_record_id, sr.filename, sr.mime_type, sr.storage_path, sr.uploaded_at,
                   sr.builtin_pack_id, sr.builtin_pack_version, sr.meta_json
            FROM knowledge_unit ku
            JOIN source_record sr ON sr.id = ku.source_record_id
            WHERE ku.id = ?
              AND sr.archived_at IS NULL
              AND (ku.archived_at IS NULL OR ku.archived_at = '')
            """,
            (unit_id,),
        ).fetchone()
    finally:
        conn.close()
    if not row:
        return None
    item = dict(row)
    item["meta"] = json.loads(item.pop("meta_json") or "{}")
    return item


def extract_chunks(path: Path, filename: str, source_type: str, mime_type: str | None) -> tuple[list[dict], int, str, dict]:
    lowered = filename.lower()
    if lowered.endswith(".docx") or mime_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ):
        return extract_docx(path, filename)
    raw = path.read_bytes()
    if source_type == "document":
        if lowered.endswith((".html", ".htm", ".xhtml", ".xml")):
            text_raw = decode_text(raw)
            text = clean_display_text(strip_html_to_text(text_raw))
            chunks = split_text(text)
            units = [
                {
                    "title": f"{filename} · 段落 {idx + 1}",
                    "content": chunk,
                    "locator": f"段落 {idx + 1}",
                    "meta": {"format": "html", "html_stripped": True, "section_path": []},
                }
                for idx, chunk in enumerate(chunks)
            ]
            note = "HTML 已剥标签、按段落切分" if units else "HTML 已导入，但未抽取到可读正文"
            return units, len(text), note, {"format": "html", "html_stripped": True}
        text = clean_display_text(decode_text(raw))
        if lowered.endswith((".md", ".markdown")):
            md_chunks = split_markdown_with_hierarchy(text)
            units = []
            for idx, c in enumerate(md_chunks):
                section_path = c.get("section_path") or []
                path_str = " · ".join(section_path) if section_path else "正文"
                units.append(
                    {
                        "title": f"{filename} · {path_str} · 段 {idx + 1}",
                        "content": c["content"],
                        "locator": f"{path_str} / 段 {idx + 1}",
                        "meta": {"format": "markdown", "section_path": section_path},
                    }
                )
            note = "Markdown 正文已按章节切分" if units else "Markdown 文件已导入，但未提取到可检索内容"
            return units, len(text), note, {"format": "markdown"}
        chunks = split_text(text)
        units = [
            {
                "title": f"{filename} · 段落 {idx + 1}",
                "content": chunk,
                "locator": f"段落 {idx + 1}",
                "meta": {"format": path.suffix.lower().lstrip(".") or "text", "section_path": []},
            }
            for idx, chunk in enumerate(chunks)
        ]
        note = "文本文件已解析" if units else "文本文件已导入，但未提取到可检索内容"
        return units, len(text), note, {"format": path.suffix.lower().lstrip(".") or "text"}

    if source_type == "spreadsheet":
        try:
            return extract_xlsx(path, filename)
        except Exception as exc:
            units = [
                {
                    "title": f"{filename} · Excel 资料",
                    "content": f"Excel 已上传，但解析失败：{exc}",
                    "locator": "整份表格",
                    "meta": {"format": "xlsx", "sheet_extract": False, "error": str(exc)},
                }
            ]
            return units, 0, "Excel 已导入，但解析失败", {"format": "xlsx", "sheet_extract": False}

    if source_type == "pdf":
        if PdfReader is None:
            units = [
                {
                    "title": f"{filename} · PDF 资料",
                    "content": "PDF 已上传，但当前环境未安装 PDF 解析能力，只保留资料记录。",
                    "locator": "整份 PDF",
                    "meta": {"format": "pdf", "lightweight_mode": True, "pdf_extract": False},
                }
            ]
            return units, 0, "PDF 已导入，但未启用正文抽取", {"format": "pdf", "pdf_extract": False}

        page_chunks: list[dict] = []
        extracted_text_parts: list[str] = []
        try:
            reader = PdfReader(str(path))
            for page_index, page in enumerate(reader.pages):
                page_text = clean_display_text(page.extract_text() or "")
                if not page_text:
                    continue
                extracted_text_parts.append(page_text)
                for chunk_index, chunk in enumerate(split_text(page_text)):
                    page_chunks.append(
                        {
                            "title": f"{filename} · 第 {page_index + 1} 页 · 片段 {chunk_index + 1}",
                            "content": chunk,
                            "locator": f"第 {page_index + 1} 页 / 片段 {chunk_index + 1}",
                            "meta": {"format": "pdf", "page": page_index + 1, "pdf_extract": True},
                        }
                    )
        except Exception as exc:
            units = [
                {
                    "title": f"{filename} · PDF 资料",
                    "content": f"PDF 已上传，但正文抽取失败：{exc}",
                    "locator": "整份 PDF",
                    "meta": {"format": "pdf", "pdf_extract": False, "error": str(exc)},
                }
            ]
            return units, 0, "PDF 已导入，但正文抽取失败", {"format": "pdf", "pdf_extract": False}

        if not page_chunks:
            units = [
                {
                    "title": f"{filename} · PDF 资料",
                    "content": "PDF 已上传，但没有抽取到可检索文字。可能是扫描型 PDF 或内容为空。",
                    "locator": "整份 PDF",
                    "meta": {"format": "pdf", "pdf_extract": False, "empty_text": True},
                }
            ]
            return units, 0, "PDF 已导入，但未抽取到正文", {"format": "pdf", "pdf_extract": False}

        return page_chunks, sum(len(t) for t in extracted_text_parts), "PDF 正文已解析", {"format": "pdf", "pdf_extract": True}

    if source_type == "image":
        fmt = path.suffix.lower().lstrip(".") or "image"
        if Image is None or pytesseract is None:
            units = [
                {
                    "title": f"{filename} · 图片资料",
                    "content": "图片已上传，但当前环境未启用 OCR，只保留资料记录。",
                    "locator": "整张图片",
                    "meta": {"format": fmt, "ocr": False, "lightweight_mode": True},
                }
            ]
            return units, 0, "图片已导入，但未启用 OCR", {"format": fmt, "ocr": False}

        try:
            image = Image.open(path)
            text = clean_display_text(pytesseract.image_to_string(image) or "")
        except Exception as exc:
            units = [
                {
                    "title": f"{filename} · 图片资料",
                    "content": f"图片已上传，但 OCR 失败：{exc}",
                    "locator": "整张图片",
                    "meta": {"format": fmt, "ocr": False, "error": str(exc)},
                }
            ]
            return units, 0, "图片已导入，但 OCR 失败", {"format": fmt, "ocr": False}

        if not text:
            units = [
                {
                    "title": f"{filename} · 图片资料",
                    "content": "图片已上传，但没有识别出明显文字。可作为图片资料保存在知识库中。",
                    "locator": "整张图片",
                    "meta": {"format": fmt, "ocr": False, "empty_text": True},
                }
            ]
            return units, 0, "图片已导入，但未识别到文字", {"format": fmt, "ocr": False}

        chunks = split_text(text)
        units = [
            {
                "title": f"{filename} · OCR 片段 {idx + 1}",
                "content": chunk,
                "locator": f"整图 OCR / 片段 {idx + 1}",
                "meta": {"format": fmt, "ocr": True},
            }
            for idx, chunk in enumerate(chunks)
        ]
        return units, len(text), "图片 OCR 已完成", {"format": fmt, "ocr": True}

    return [], 0, "未知文件类型", {}


class Handler(BaseHTTPRequestHandler):
    server_version = "KBMVP/0.1"

    def _send_json(self, status: int, payload: dict):
        raw = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "*")
        self.end_headers()
        self.wfile.write(raw)

    def _send_file(self, path: Path, content_type: str = "text/html; charset=utf-8"):
        raw = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.end_headers()
        self.wfile.write(raw)

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "*")
        self.end_headers()

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/health":
            return self._send_json(200, {
                "status": "ok",
                "llm_enabled": LLM_ENABLED,
                "embedding_enabled": EMBEDDING_ENABLED,
                "embedding_key_configured": EMBEDDING_KEY_CONFIGURED,
                "embedding_env_forced_off": EMBEDDING_ENV_FORCED_OFF,
            })
        if parsed.path in ("/", "/demo.html"):
            return self._send_file(DEMO_HTML)
        if parsed.path == "/api/v1/admin/ingestion-jobs":
            limit = int(parse_qs(parsed.query).get("limit", ["20"])[0])
            conn = get_conn()
            try:
                rows = conn.execute(
                    "SELECT * FROM ingestion_job ORDER BY created_at DESC LIMIT ?",
                    (limit,),
                ).fetchall()
                items = [dict(r) for r in rows]
            finally:
                conn.close()
            return self._send_json(200, {"items": items})
        if parsed.path.startswith("/api/v1/admin/ingestion-jobs/") and parsed.path.endswith("/progress"):
            job_id = parsed.path[len("/api/v1/admin/ingestion-jobs/"):-len("/progress")]
            if not job_id:
                return self._send_json(400, {"detail": "missing job id"})
            conn = get_conn()
            try:
                row = conn.execute(
                    """
                    SELECT id, filename, source_type, status, created_at, updated_at, finished_at,
                           progress_pct, current_stage, unit_count, note
                    FROM ingestion_job WHERE id=?
                    """,
                    (job_id,),
                ).fetchone()
            finally:
                conn.close()
            if not row:
                return self._send_json(404, {"detail": "job not found"})
            return self._send_json(200, dict(row))
        if parsed.path == "/api/v1/admin/units":
            params = parse_qs(parsed.query)
            limit = int(params.get("limit", ["50"])[0])
            include_archived = parse_bool(params.get("include_archived", [None])[0], default=False)
            filters = normalize_filters(
                {
                    "source_scope": params.get("source_scope", [""])[0],
                    "source_type": params.get("source_type", [""])[0],
                    "builtin_pack_id": params.get("builtin_pack_id", [""])[0],
                    "filename": params.get("filename", [""])[0],
                }
            )
            where = []
            query_params: list[object] = []
            if not include_archived:
                where.append("sr.archived_at IS NULL")
                where.append("(ku.archived_at IS NULL OR ku.archived_at = '')")
            if filters.get("source_scope"):
                where.append("sr.source_scope = ?")
                query_params.append(filters["source_scope"])
            if filters.get("source_type"):
                where.append("ku.source_type = ?")
                query_params.append(filters["source_type"])
            if filters.get("builtin_pack_id"):
                where.append("sr.builtin_pack_id = ?")
                query_params.append(filters["builtin_pack_id"])
            if filters.get("filename"):
                where.append("sr.filename LIKE ?")
                query_params.append(f"%{filters['filename']}%")
            where_clause = f"WHERE {' AND '.join(where)}" if where else ""
            conn = get_conn()
            try:
                rows = conn.execute(
                    f"""
                    SELECT ku.id, ku.source_type, ku.source_scope, ku.title, ku.content, ku.locator, ku.created_at,
                           sr.filename, sr.uploaded_at, sr.builtin_pack_id, sr.builtin_pack_version, sr.meta_json
                    FROM knowledge_unit ku
                    JOIN source_record sr ON sr.id = ku.source_record_id
                    {where_clause}
                    ORDER BY ku.created_at DESC, ku.chunk_index ASC
                    LIMIT ?
                    """,
                    [*query_params, limit],
                ).fetchall()
                items = []
                for row in rows:
                    item = dict(row)
                    item["meta"] = json.loads(item.pop("meta_json") or "{}")
                    items.append(item)
            finally:
                conn.close()
            return self._send_json(200, {"items": items})
        if parsed.path == "/api/v1/admin/source-records":
            params = parse_qs(parsed.query)
            limit = int(params.get("limit", ["50"])[0])
            offset = int(params.get("offset", ["0"])[0])
            include_archived = parse_bool(params.get("include_archived", [None])[0], default=False)
            filters = {
                "source_scope": params.get("source_scope", [""])[0],
                "source_type": params.get("source_type", [""])[0],
                "builtin_pack_id": params.get("builtin_pack_id", [""])[0],
                "filename": params.get("filename", [""])[0],
            }
            return self._send_json(
                200,
                list_source_records(
                    limit=limit,
                    offset=offset,
                    filters=filters,
                    include_archived=include_archived,
                ),
            )
        if parsed.path == "/api/v1/admin/source-records/detail":
            source_record_id = parse_qs(parsed.query).get("id", [""])[0]
            include_archived = parse_bool(parse_qs(parsed.query).get("include_archived", [None])[0], default=False)
            if not source_record_id:
                return self._send_json(400, {"detail": "missing id"})
            item = get_source_record_detail(int(source_record_id), include_archived=include_archived)
            if not item:
                return self._send_json(404, {"detail": "source record not found"})
            return self._send_json(200, item)
        if parsed.path == "/api/v1/admin/source-summary":
            return self._send_json(200, {"items": list_source_summary()})
        if parsed.path == "/api/v1/admin/builtin-knowledge/packs":
            return self._send_json(200, {"items": list_builtin_pack_status()})
        if parsed.path == "/api/v1/knowledge/unit":
            unit_id = parse_qs(parsed.query).get("id", [""])[0]
            if not unit_id:
                return self._send_json(400, {"detail": "missing id"})
            item = get_unit_detail(unit_id)
            if not item:
                return self._send_json(404, {"detail": "unit not found"})
            storage_path = item.get("storage_path") or ""
            file_url = (
                f"/api/v1/knowledge/file?id={unit_id}"
                if storage_path and Path(storage_path).exists()
                else None
            )
            return self._send_json(
                200,
                {
                    "id": item["id"],
                    "title": clean_display_text(item["title"]),
                    "content": clean_display_text(item["content"]),
                    "locator": item["locator"],
                    "source_type": normalize_response_source_type(item["source_type"]),
                    "source_scope": item["source_scope"],
                    "filename": item["filename"],
                    "mime_type": item["mime_type"],
                    "uploaded_at": item["uploaded_at"],
                    "builtin_pack_id": item["builtin_pack_id"],
                    "builtin_pack_version": item["builtin_pack_version"],
                    "meta": item["meta"],
                    "file_url": file_url,
                },
            )
        if parsed.path == "/api/v1/knowledge/file":
            unit_id = parse_qs(parsed.query).get("id", [""])[0]
            if not unit_id:
                return self._send_json(400, {"detail": "missing id"})
            item = get_unit_detail(unit_id)
            if not item:
                return self._send_json(404, {"detail": "unit not found"})
            path = Path(item["storage_path"])
            if not path.exists():
                return self._send_json(404, {"detail": "file not found"})
            mime = item["mime_type"] or "application/octet-stream"
            return self._send_file(path, mime)
        return self._send_json(404, {"detail": "not found"})

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/v1/knowledge/query":
            try:
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length)
                payload = json.loads(raw.decode("utf-8") or "{}")
            except Exception:
                return self._send_json(400, {"detail": "invalid json"})
            query = (payload.get("query") or "").strip()
            if not query:
                return self._send_json(400, {"detail": "missing query"})
            return self._send_json(200, build_query_response(query, payload.get("filters")))
        if parsed.path == "/api/v1/admin/embedding/toggle":
            try:
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length)
                payload = json.loads(raw.decode("utf-8") or "{}")
            except Exception:
                return self._send_json(400, {"detail": "invalid json"})
            on = bool(payload.get("enabled"))
            ok, reason = toggle_embedding(on)
            return self._send_json(
                200 if ok else 409,
                {**embedding_status(), "changed": ok, "reject_reason": None if ok else reason},
            )
        if parsed.path == "/api/v1/admin/embeddings/reindex":
            params = parse_qs(parsed.query)
            force = params.get("force", ["false"])[0].lower() == "true"
            conn = get_conn()
            try:
                if force:
                    rows = conn.execute(
                        "SELECT ku.id FROM knowledge_unit ku JOIN source_record sr ON sr.id=ku.source_record_id "
                        "WHERE sr.archived_at IS NULL AND (ku.archived_at IS NULL OR ku.archived_at='')"
                    ).fetchall()
                else:
                    rows = conn.execute(
                        "SELECT ku.id FROM knowledge_unit ku JOIN source_record sr ON sr.id=ku.source_record_id "
                        "WHERE ku.embedding IS NULL AND sr.archived_at IS NULL AND (ku.archived_at IS NULL OR ku.archived_at='')"
                    ).fetchall()
            finally:
                conn.close()
            ku_ids = [r["id"] for r in rows]
            if not EMBEDDING_ENABLED:
                return self._send_json(503, {
                    "detail": "embedding disabled",
                    "reason": "disabled",
                    "hint": "export DASHSCOPE_API_KEY after signing up for 阿里云百炼, then restart",
                })
            updated = compute_unit_embeddings(ku_ids)
            return self._send_json(200, {
                "requested": len(ku_ids),
                "embedded": updated,
                "force": force,
            })
        if parsed.path == "/api/v1/chat/rag-synthesize":
            try:
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length)
                payload = json.loads(raw.decode("utf-8") or "{}")
            except Exception:
                return self._send_json(400, {"detail": "invalid json"})
            query = (payload.get("query") or "").strip()
            evidence_ids = payload.get("evidence_ids") or []
            if not query or not evidence_ids:
                return self._send_json(400, {"detail": "missing query or evidence_ids"})
            if not LLM_ENABLED and not (payload.get("api_key") or "").strip():
                return self._send_json(503, {
                    "detail": "llm unavailable",
                    "reason": "disabled",
                    "hint": "需要配置 DEEPSEEK_API_KEY 才能合成答案",
                })
            conn = get_conn()
            try:
                placeholders = ",".join(["?"] * len(evidence_ids))
                crows = conn.execute(
                    f"SELECT ku.id, ku.title, ku.content, ku.locator, sr.filename "
                    f"FROM knowledge_unit ku JOIN source_record sr ON sr.id=ku.source_record_id "
                    f"WHERE ku.id IN ({placeholders})",
                    list(evidence_ids),
                ).fetchall()
            finally:
                conn.close()
            if not crows:
                return self._send_json(404, {"detail": "no chunks found"})
            chunk_blocks: list[str] = []
            ordered_ids: list[str] = []
            for idx, r in enumerate(crows, start=1):
                ordered_ids.append(r["id"])
                title = (r["title"] or "").strip()
                content = (r["content"] or "").strip()
                locator = (r["locator"] or "").strip()
                src = (r["filename"] or "").strip()
                block = f"[{idx}] {title}\n来源: {src} · 定位: {locator}\n{content[:1500]}"
                chunk_blocks.append(block)
            user_content = "资料:\n\n" + "\n\n".join(chunk_blocks) + f"\n\n用户问题: {query}"
            messages = [
                {"role": "system", "content": RAG_SYSTEM_PROMPT},
                {"role": "user", "content": user_content},
            ]
            client_api_key = (payload.get("api_key") or "").strip() or None
            request_id = f"rag_{uuid.uuid4().hex[:10]}"
            try:
                result = llm_provider.call_llm(
                    LLM_DEFAULT_PROVIDER,
                    messages,
                    timeout_s=30.0,
                    request_id=request_id,
                    api_key=client_api_key,
                )
            except llm_provider.LLMTimeout as exc:
                return self._send_json(503, {"detail": "llm timeout", "reason": "timeout", "error": str(exc)})
            except llm_provider.LLMRateLimit as exc:
                return self._send_json(503, {"detail": "llm rate limited", "reason": "rate_limited", "error": str(exc)})
            except llm_provider.LLMInvalidResponse as exc:
                return self._send_json(503, {"detail": "llm invalid response", "reason": "invalid_response", "error": str(exc)})
            except llm_provider.LLMError as exc:
                return self._send_json(503, {"detail": "llm error", "reason": getattr(exc, "reason", "unknown"), "error": str(exc)})
            return self._send_json(200, {
                **result,
                "evidence_ids_ordered": ordered_ids,
            })
        if parsed.path == "/api/v1/chat/llm-fallback":
            try:
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length)
                payload = json.loads(raw.decode("utf-8") or "{}")
            except Exception:
                return self._send_json(400, {"detail": "invalid json"})
            query = (payload.get("query") or "").strip()
            if not query:
                return self._send_json(400, {"detail": "missing query"})
            provider = (payload.get("provider") or LLM_DEFAULT_PROVIDER).strip()
            model = payload.get("model")
            client_api_key = (payload.get("api_key") or "").strip() or None
            request_id = f"req_{uuid.uuid4().hex[:10]}"
            messages = [
                {"role": "system", "content": LLM_FALLBACK_SYSTEM_PROMPT},
                {"role": "user", "content": query},
            ]
            try:
                result = llm_provider.call_llm(
                    provider,
                    messages,
                    model=model,
                    timeout_s=30.0,
                    request_id=request_id,
                    api_key=client_api_key,
                )
            except llm_provider.LLMDisabled as exc:
                return self._send_json(
                    503,
                    {
                        "detail": "llm provider unavailable",
                        "reason": "disabled",
                        "error": str(exc),
                        "hint": "未配置 key：在页面「配置启明」里粘贴 key 保存到本浏览器，或服务器端 export DEEPSEEK_API_KEY 后重启",
                    },
                )
            except llm_provider.LLMTimeout as exc:
                return self._send_json(
                    503,
                    {"detail": "llm timeout", "reason": "timeout", "error": str(exc)},
                )
            except llm_provider.LLMRateLimit as exc:
                return self._send_json(
                    503,
                    {"detail": "llm rate limited", "reason": "rate_limited", "error": str(exc)},
                )
            except llm_provider.LLMInvalidResponse as exc:
                return self._send_json(
                    503,
                    {"detail": "llm invalid response", "reason": "invalid_response", "error": str(exc)},
                )
            except llm_provider.LLMProviderError as exc:
                return self._send_json(
                    503,
                    {"detail": "llm provider error", "reason": "server_error", "error": str(exc)},
                )
            return self._send_json(200, result)
        if parsed.path == "/api/v1/knowledge/llm-capture":
            try:
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length)
                payload = json.loads(raw.decode("utf-8") or "{}")
            except Exception:
                return self._send_json(400, {"detail": "invalid json"})
            query = (payload.get("query") or "").strip()
            answer = (payload.get("answer") or "").strip()
            edited = (payload.get("edited_answer") or "").strip()
            content = edited or answer
            if not query or not content:
                return self._send_json(400, {"detail": "missing query or answer"})
            if len(content.encode("utf-8")) > MAX_MANUAL_ENTRY_BYTES:
                return self._send_json(
                    413,
                    {"detail": f"内容过长（上限 {MAX_MANUAL_ENTRY_BYTES // 1024}KB）"},
                )
            title_override = (payload.get("title_override") or "").strip()
            title = title_override or f"{query[:40]} · AI 兜底"
            meta = {
                "llm_generated": True,
                "llm_model": (payload.get("llm_model") or "deepseek-chat"),
                "llm_provider": (payload.get("llm_provider") or LLM_DEFAULT_PROVIDER),
                "user_edited": bool(edited and edited != answer),
                "source_query": query,
                "request_id": payload.get("request_id"),
                "scope_label": "AI 生成 · 运行时沉淀",
                "original_answer": answer if (edited and edited != answer) else None,
            }
            record = insert_curated_knowledge(
                title=clean_display_text(title),
                content=clean_display_text(content),
                meta=meta,
                source_scope="runtime_curated",
                source_type="document",
                note="AI 兜底沉淀",
            )
            print(
                f"[curate] {json.dumps({'ts': now_iso(), 'endpoint': 'llm-capture', 'source_record_id': record['source_record_id'], 'content_length': len(content), 'user_edited': meta['user_edited']}, ensure_ascii=False)}",
                flush=True,
            )
            return self._send_json(200, {**record, "meta": meta})
        if parsed.path == "/api/v1/knowledge/manual-entry":
            try:
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length)
                payload = json.loads(raw.decode("utf-8") or "{}")
            except Exception:
                return self._send_json(400, {"detail": "invalid json"})
            title = (payload.get("title") or "").strip()
            content = (payload.get("content") or "").strip()
            if not title or not content:
                return self._send_json(400, {"detail": "标题和内容都得填"})
            if len(title) > 120:
                return self._send_json(400, {"detail": "标题不超过 120 字符"})
            if len(content.encode("utf-8")) > MAX_MANUAL_ENTRY_BYTES:
                return self._send_json(
                    413,
                    {"detail": f"内容过长（上限 {MAX_MANUAL_ENTRY_BYTES // 1024}KB）"},
                )
            tags = payload.get("tags") or []
            if isinstance(tags, str):
                tags = [t.strip() for t in tags.split(",") if t.strip()]
            meta = {
                "manually_entered": True,
                "source_query": (payload.get("source_query") or "").strip() or None,
                "tags": list(tags) if isinstance(tags, list) else [],
                "scope_label": "运行时沉淀",
            }
            record = insert_curated_knowledge(
                title=clean_display_text(title),
                content=clean_display_text(content),
                meta=meta,
                source_scope="runtime_curated",
                source_type="document",
                note="手动录入",
            )
            print(
                f"[curate] {json.dumps({'ts': now_iso(), 'endpoint': 'manual-entry', 'source_record_id': record['source_record_id'], 'content_length': len(content)}, ensure_ascii=False)}",
                flush=True,
            )
            return self._send_json(200, {**record, "meta": meta})
        if parsed.path == "/api/v1/admin/builtin-knowledge/reload":
            params = parse_qs(parsed.query)
            pack_id = params.get("pack_id", [None])[0]
            force = params.get("force", ["true"])[0].lower() != "false"
            try:
                result = import_enabled_builtin_packs(force=force, pack_id=pack_id)
            except Exception as exc:
                return self._send_json(500, {"detail": str(exc)})
            return self._send_json(200, result)
        if parsed.path == "/api/v1/admin/source-records/archive":
            try:
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length)
                payload = json.loads(raw.decode("utf-8") or "{}")
            except Exception:
                return self._send_json(400, {"detail": "invalid json"})
            source_record_ids = payload.get("source_record_ids") or []
            reason = (payload.get("reason") or "manual archive").strip()
            result = archive_source_records(source_record_ids, reason)
            return self._send_json(200, result)
        if parsed.path == "/api/v1/admin/source-records/unarchive":
            try:
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length)
                payload = json.loads(raw.decode("utf-8") or "{}")
            except Exception:
                return self._send_json(400, {"detail": "invalid json"})
            source_record_ids = payload.get("source_record_ids") or []
            result = unarchive_source_records(source_record_ids)
            return self._send_json(200, result)
        if parsed.path == "/api/v1/admin/source-records/update":
            try:
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length)
                payload = json.loads(raw.decode("utf-8") or "{}")
            except Exception:
                return self._send_json(400, {"detail": "invalid json"})
            source_record_id = payload.get("source_record_id")
            if not source_record_id:
                return self._send_json(400, {"detail": "missing source_record_id"})
            try:
                result = update_source_record_metadata(
                    int(source_record_id),
                    display_title=payload.get("display_title"),
                    tags=payload.get("tags"),
                    note=payload.get("note"),
                    source_scope=payload.get("source_scope"),
                )
            except ValueError as exc:
                return self._send_json(400, {"detail": str(exc)})
            return self._send_json(200, result)

        if parsed.path != "/api/v1/knowledge/ingest":
            return self._send_json(404, {"detail": "not found"})

        try:
            uploaded = parse_uploaded_file(self)
        except ValueError as exc:
            return self._send_json(
                413,
                {
                    "detail": f"文件过大（上限 {MAX_UPLOAD_BYTES // (1024 * 1024)}MB）",
                    "error": str(exc),
                    "max_bytes": MAX_UPLOAD_BYTES,
                },
            )
        if uploaded is None or not uploaded.get("filename"):
            return self._send_json(400, {"detail": "missing file"})

        filename = os.path.basename(uploaded["filename"])
        raw = uploaded["raw"]
        if not raw:
            return self._send_json(400, {"detail": "empty file"})

        source_type = detect_source_type(filename, uploaded.get("mime_type"))
        mime_type = uploaded.get("mime_type")
        job_id = f"job_{uuid.uuid4().hex[:10]}"
        created_at = now_iso()

        safe_name = f"{uuid.uuid4().hex}{Path(filename).suffix}"
        target = UPLOAD_DIR / safe_name
        try:
            target.write_bytes(raw)
        except Exception as exc:
            return self._send_json(500, {"detail": f"保存上传文件失败：{exc}"})

        conn = get_conn()
        try:
            conn.execute(
                """
                INSERT INTO ingestion_job (
                    id, filename, source_type, status, created_at, updated_at,
                    progress_pct, current_stage, unit_count, note
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    job_id,
                    filename,
                    source_type,
                    "queued",
                    created_at,
                    created_at,
                    INGEST_STAGE_PCT["queued"],
                    "queued",
                    0,
                    "已接收，排队中",
                ),
            )
            conn.commit()
        finally:
            conn.close()

        worker = threading.Thread(
            target=_run_ingest_worker,
            args=(job_id, str(target), filename, source_type, mime_type),
            daemon=True,
            name=f"ingest-worker-{job_id}",
        )
        worker.start()

        return self._send_json(
            200,
            {
                "job_id": job_id,
                "filename": filename,
                "source_type": source_type,
                "status": "queued",
                "unit_count": 0,
                "preview_units": [],
                "note": "已接收，后台处理中",
                "poll_url": f"/api/v1/admin/ingestion-jobs/{job_id}/progress",
            },
        )


def main():
    init_db()
    reaped = reap_orphan_jobs()
    if reaped > 0:
        print(f"[startup] reaped {reaped} orphan ingestion job(s)", flush=True)
    if LLM_ENABLED:
        print("[llm] deepseek enabled (DEEPSEEK_API_KEY present)", flush=True)
    else:
        print("[llm] disabled (no DEEPSEEK_API_KEY in env)", flush=True)
    import_enabled_builtin_packs(force=False)
    port = int(sys.argv[1]) if len(sys.argv) > 1 else 38427
    server = ThreadingHTTPServer(("0.0.0.0", port), Handler)
    print(f"KB MVP server listening on 0.0.0.0:{port}", flush=True)
    server.serve_forever()


if __name__ == "__main__":
    main()
